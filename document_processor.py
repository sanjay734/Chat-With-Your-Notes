"""
Document Processor for Multi-Format RAG Assistant
Handles ingestion, parsing, and chunking of PDFs, DOCX, CSV, and TXT files
"""

import os
from pypdf import PdfReader
import pdfplumber
from docx import Document as DocxDocument
from typing import List, Tuple, Optional
import csv
from io import BytesIO
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Handles document ingestion and processing for multiple formats"""
    
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
    
    def process_file(self, file_path: str, file_type: str) -> List[dict]:
        """
        Process a file based on its type and return chunks with metadata
        
        Args:
            file_path: Path to the file
            file_type: Type of file ('pdf', 'docx', 'csv', 'txt')
            
        Returns:
            List of dictionaries containing chunk text and metadata
        """
        logger.info(f"Processing {file_type} file: {file_path}")
        
        if file_type.lower() == 'pdf':
            return self._process_pdf(file_path)
        elif file_type.lower() == 'docx':
            return self._process_docx(file_path)
        elif file_type.lower() == 'csv':
            return self._process_csv(file_path)
        elif file_type.lower() == 'txt':
            return self._process_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
    
    def _process_pdf(self, file_path: str) -> List[dict]:
        """Extract text from PDF with page-level metadata"""
        chunks = []
        
        try:
            # Try with pdfplumber first (better layout handling)
            with pdfplumber.open(file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text()
                    
                    if text.strip():
                        # Also try to extract tables
                        tables = page.extract_tables()
                        table_text = ""
                        if tables:
                            for table in tables:
                                table_text += "\n[TABLE]\n"
                                for row in table:
                                    table_text += " | ".join(str(cell) if cell else "" for cell in row) + "\n"
                        
                        full_text = text + table_text
                        
                        # Chunk the page text
                        page_chunks = self._chunk_text(full_text)
                        
                        for chunk_num, chunk in enumerate(page_chunks):
                            chunks.append({
                                'content': chunk,
                                'metadata': {
                                    'source': os.path.basename(file_path),
                                    'file_type': 'pdf',
                                    'page': page_num,
                                    'chunk': chunk_num
                                }
                            })
            
            logger.info(f"PDF processed: {len(chunks)} chunks created")
            return chunks
            
        except Exception as e:
            logger.error(f"Error processing PDF with pdfplumber: {e}")
            # Fallback to PyPDF2
            return self._process_pdf_pypdf2(file_path)
    
    def _process_pdf_pypdf2(self, file_path: str) -> List[dict]:
        """Fallback PDF processing with pypdf"""
        chunks = []
        
        with open(file_path, 'rb') as file:
            reader = PdfReader(file)
            
            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                
                if text.strip():
                    page_chunks = self._chunk_text(text)
                    
                    for chunk_num, chunk in enumerate(page_chunks):
                        chunks.append({
                            'content': chunk,
                            'metadata': {
                                'source': os.path.basename(file_path),
                                'file_type': 'pdf',
                                'page': page_num,
                                'chunk': chunk_num
                            }
                        })
        
        logger.info(f"PDF processed (pypdf fallback): {len(chunks)} chunks created")
        return chunks
    
    def _process_docx(self, file_path: str) -> List[dict]:
        """Extract text from Word documents"""
        chunks = []
        doc = DocxDocument(file_path)
        
        full_text = ""
        paragraph_metadata = []
        
        for para_num, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                full_text += paragraph.text + "\n"
                paragraph_metadata.append((len(full_text), para_num))
        
        # Also extract from tables
        for table in doc.tables:
            full_text += "\n[TABLE]\n"
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                full_text += row_text + "\n"
        
        # Chunk the text
        page_chunks = self._chunk_text(full_text)
        
        for chunk_num, chunk in enumerate(page_chunks):
            chunks.append({
                'content': chunk,
                'metadata': {
                    'source': os.path.basename(file_path),
                    'file_type': 'docx',
                    'chunk': chunk_num
                }
            })
        
        logger.info(f"DOCX processed: {len(chunks)} chunks created")
        return chunks
    
    def _process_csv(self, file_path: str) -> List[dict]:
        """Process CSV files"""
        chunks = []
        
        with open(file_path, 'r', encoding='utf-8') as file:
            reader = csv.DictReader(file)
            
            # Convert to formatted text
            full_text = ""
            for row_num, row in enumerate(reader, 1):
                row_text = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
                full_text += row_text + "\n"
        
        # Chunk the text
        page_chunks = self._chunk_text(full_text)
        
        for chunk_num, chunk in enumerate(page_chunks):
            chunks.append({
                'content': chunk,
                'metadata': {
                    'source': os.path.basename(file_path),
                    'file_type': 'csv',
                    'chunk': chunk_num
                }
            })
        
        logger.info(f"CSV processed: {len(chunks)} chunks created")
        return chunks
    
    def _process_txt(self, file_path: str) -> List[dict]:
        """Process plain text files"""
        chunks = []
        
        with open(file_path, 'r', encoding='utf-8') as file:
            full_text = file.read()
        
        # Chunk the text
        page_chunks = self._chunk_text(full_text)
        
        for chunk_num, chunk in enumerate(page_chunks):
            chunks.append({
                'content': chunk,
                'metadata': {
                    'source': os.path.basename(file_path),
                    'file_type': 'txt',
                    'chunk': chunk_num
                }
            })
        
        logger.info(f"TXT processed: {len(chunks)} chunks created")
        return chunks
    
    def _chunk_text(self, text: str, chunk_size: int = None, overlap: int = None) -> List[str]:
        """
        Split text into overlapping chunks
        
        Args:
            text: Text to chunk
            chunk_size: Size of each chunk
            overlap: Overlap between chunks
            
        Returns:
            List of text chunks
        """
        chunk_size = chunk_size or self.chunk_size
        overlap = overlap or self.chunk_overlap
        
        if not text or len(text) < chunk_size:
            return [text] if text.strip() else []
        
        chunks = []
        start = 0
        text_len = len(text)
        
        while start < text_len:
            end = min(start + chunk_size, text_len)
            chunk = text[start:end]
            is_final_chunk = (end >= text_len)
            
            # Try to break at sentence boundary (only when there's more text after this chunk)
            if not is_final_chunk:
                last_period = chunk.rfind('.')
                last_newline = chunk.rfind('\n')
                break_point = max(last_period, last_newline)
                
                if break_point > chunk_size // 2:
                    end = start + break_point + 1
            
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            # Reached the end of the text — stop, otherwise the overlap
            # subtraction on a short final slice can push start backward
            # and loop forever.
            if is_final_chunk:
                break
            
            next_start = end - overlap
            # Guard against non-forward progress (e.g. overlap >= chunk span)
            start = next_start if next_start > start else start + 1
        
        return chunks
    
    def process_uploaded_file(self, uploaded_file) -> Tuple[List[dict], str]:
        """
        Process an uploaded file from Streamlit
        
        Args:
            uploaded_file: Streamlit uploaded file object
            
        Returns:
            Tuple of (chunks list, file extension)
        """
        # Save uploaded file temporarily
        file_extension = uploaded_file.name.split('.')[-1].lower()
        temp_path = f"./temp_{uploaded_file.name}"
        
        with open(temp_path, 'wb') as f:
            f.write(uploaded_file.getbuffer())
        
        try:
            chunks = self.process_file(temp_path, file_extension)
            # Show the real file name in sources, not the temp file name
            for chunk in chunks:
                chunk['metadata']['source'] = uploaded_file.name
            return chunks, file_extension
        finally:
            # Clean up temp file
            if os.path.exists(temp_path):
                os.remove(temp_path)


def validate_file_type(file_name: str) -> bool:
    """Check if file type is supported"""
    supported = ['.pdf', '.docx', '.csv', '.txt']
    return any(file_name.lower().endswith(ext) for ext in supported)
